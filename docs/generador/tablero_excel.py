# Genera el Tablero de Indicadores CENTEC (Excel) con fórmulas vivas. Uso: python3 tablero_excel.py salida.xlsx
import sys, json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.properties import CalcProperties

OUT = sys.argv[1] if len(sys.argv) > 1 else 'Tablero_Indicadores_CENTEC.xlsx'
IND = json.load(open('../../data/indicadores.json'))
NAVY, ORANGE = '1B2A6B', 'D9771E'
F = lambda **k: Font(name='Arial', size=10, **k)
H = Font(name='Arial', size=10, bold=True, color='FFFFFF')
HF = PatternFill('solid', fgColor=NAVY)
INP = PatternFill('solid', fgColor='FFF7CC')
thin = Side(style='thin', color='C8CED8'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
wb = Workbook()

def head(ws, row, vals, widths=None):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row, i, v); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center'); c.border = B
    if widths:
        for i, w in enumerate(widths, 1): ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[row].height = 54

# ---------- Instrucciones ----------
ws = wb.active; ws.title = 'Instrucciones'
ws.column_dimensions['A'].width = 110
txt = [
 ('TABLERO DE INDICADORES — PLANEACIÓN ACADÉMICA CENTEC (multi-año: 2027, 2028, …)', Font(name='Arial', size=14, bold=True, color=NAVY)),
 ('Complemento en Excel del aplicativo Sistema de Planeación CENTEC. Sirve para auditar en papel/Excel o cuando no haya conexión. Las fórmulas son las mismas del aplicativo (v%s).' % IND['version'], F()),
 ('', F()),
 ('CÓMO USARLO', F(bold=True, color=ORANGE)),
 ('1. Celdas con fondo amarillo = datos que usted digita. Celdas blancas = fórmulas (no las edite).', F()),
 ('2. Hoja «Registro»: una fila por docente y período, para cualquier año (columna Año). Las filas marcadas EJEMPLO muestran el formato: bórrelas antes de usar.', F()),
 ('3. Hoja «Parámetros»: pesos, umbrales y constantes. Si se cambian, suba la versión: el comparativo año a año dejará de ser comparable.', F()),
 ('4. Hoja «Tablero»: resumen por docente y por año. Hoja «Año vs año»: compara automáticamente cada indicador entre 2027 y 2031.', F()),
 ('5. Hoja «Lecciones»: decisiones del cierre anual (mantener / ajustar / eliminar / escalar) que alimentan la planeación del año siguiente.', F()),
 ('', F()),
 ('REGLAS QUE EVITAN EL SESGO', F(bold=True, color=ORANGE)),
] + [('• ' + r, F()) for r in IND['reglas_bloqueo']] + [
 ('• La normalización por observador (indulgencia/severidad) la calcula el aplicativo; en Excel digite el promedio Teach ya normalizado que reporta el aplicativo, o el promedio crudo si no lo usa.', F()),
 ('', F()),
 ('FUENTES: Teach (Banco Mundial) para observación; ganancia normalizada de Hake (1998) con contracción bayesiana; MET Project para pesos; EEF Toolkit y Dunlosky et al. (2013) para prácticas.', F(italic=True)),
]
for i, (t, f) in enumerate(txt, 1):
    c = ws.cell(i, 1, t); c.font = f; c.alignment = Alignment(wrap_text=True)

# ---------- Parámetros ----------
P = wb.create_sheet('Parámetros')
head(P, 1, ['Parámetro', 'Valor', 'Descripción / fuente'], [34, 12, 90])
params = [(c['codigo'] + ' ' + c['nombre'], c['peso'], c['formula']) for c in IND['componentes']] + [
 ('Umbral verde', 85, 'ICD ≥ 85'), ('Umbral amarillo', 70, 'ICD 70–84; menor que 70 = rojo'),
 ('Kappa mínimo', 0.60, 'Landis y Koch: ≥ 0,61 sustancial. Por debajo, la observación no cuenta'),
 ('k de contracción', 10, 'g ajustada = n/(n+k)·g + k/(n+k)·g media'),
 ('g media del colegio', 0.30, 'Actualizar cada año con el promedio institucional real (2027: supuesto inicial, Hake 1998)'),
 ('Mínimo de fuentes', 3, 'Sin 3 fuentes no se asigna color'),
 ('Versión de fórmulas', IND['version'], 'Cambia si cambia cualquier peso o fórmula')]
for i, (a, b, c) in enumerate(params, 2):
    P.cell(i, 1, a).font = F(); x = P.cell(i, 2, b); x.font = F(color='0000FF'); x.fill = INP; P.cell(i, 3, c).font = F()
    for j in (1, 2, 3): P.cell(i, j).border = B
    if i <= 11: x.number_format = '0%'
P.cell(12, 1, 'Suma de pesos (debe ser 100 %)').font = F(bold=True); P['B12'] = '=SUM(B2:B11)'; P['B12'].number_format = '0%'; P['B12'].font = F(bold=True)
# desplazar parámetros posteriores una fila: reescribir desde fila 13
for i, (a, b, c) in enumerate(params[10:], 13):
    P.cell(i, 1, a).font = F(); x = P.cell(i, 2, b); x.font = F(color='0000FF'); x.fill = INP; P.cell(i, 3, c).font = F()
    for j in (1, 2, 3): P.cell(i, j).border = B
P['B12'].fill = PatternFill(fill_type=None); P['A12'].fill = PatternFill(fill_type=None)
P['C12'] = 'Control'; P['C12'].font = F(italic=True)
W = {k: f"Parámetros!$B${i}" for i, k in enumerate(['C1','C2','C3','C4','C5','C6','C7','C8','C9','C10'], 2)}
VERDE, AMAR, KMIN, KC, GM, MINF = 'Parámetros!$B$13', 'Parámetros!$B$14', 'Parámetros!$B$15', 'Parámetros!$B$16', 'Parámetros!$B$17', 'Parámetros!$B$18'

# ---------- Registro ----------
R = wb.create_sheet('Registro')
cols = [
 ('Año', 7, 'in'), ('Período', 7, 'in'), ('Docente', 24, 'in'), ('Grado / área', 16, 'in'),
 ('Planeadores exigidos', 10, 'in'), ('Planeadores a tiempo', 10, 'in'), ('Semanas planeadas', 9, 'in'), ('Semanas con evidencia', 10, 'in'),
 ('DBA programados', 9, 'in'), ('DBA evidenciados', 10, 'in'), ('R1 total (0–24)', 9, 'in'), ('R2 promedio (0–3)', 9, 'in'),
 ('Teach promedio (1–5)', 9, 'in'), ('κ ponderado (vacío si 1 obs.)', 10, 'in'), ('Observadores distintos', 10, 'in'),
 ('g promedio estudiantes', 10, 'in'), ('n estudiantes (pre ≤ 90)', 10, 'in'), ('Ajustes PIAR comprometidos', 10, 'in'), ('Ajustes PIAR verificados', 10, 'in'),
 ('Fases transversales evidenciadas', 10, 'in'), ('Proyectos asignados', 9, 'in'), ('Encuesta promedio (1–5)', 9, 'in'), ('n encuesta', 8, 'in'),
 ('C1 Puntualidad', 9, 'f'), ('C2 Ejecución', 9, 'f'), ('C3 Cobertura DBA', 9, 'f'), ('C4 Planeador', 9, 'f'), ('C5 Cuadernos', 9, 'f'),
 ('C6 Observación', 9, 'f'), ('C7 Aprendizaje', 9, 'f'), ('C8 PIAR', 9, 'f'), ('C9 Transversal', 9, 'f'), ('C10 Estudiantes', 9, 'f'),
 ('Fuentes con dato', 8, 'f'), ('ICD sin tope', 9, 'f'), ('Tope por reglas', 9, 'f'), ('ICD', 9, 'f'), ('Estado', 20, 'f'), ('Regla aplicada', 40, 'f'),
]
head(R, 1, [c[0] for c in cols], [c[1] for c in cols])
R.freeze_panes = 'E2'
import random
random.seed(2027)
ej = []
_doc = [('EJEMPLO Docente A','Transición',0.92),('EJEMPLO Docente B','1° Integrada',0.78),('EJEMPLO Docente C','3° Integrada',0.62),('EJEMPLO Docente D','5° Integrada',0.88),
        ('EJEMPLO Docente E','Inglés primaria',0.72),('EJEMPLO Docente F','Matemáticas 6°–9°',0.48),('EJEMPLO Docente G','Lenguaje 6°–11°',0.84),('EJEMPLO Docente H','Ciencias 6°–11°',0.68)]
for anio, mej in ((2027, 0), (2028, 0.06)):
    for per in ((1, 2) if anio == 2027 else (1,)):
        for nom, gr, p in _doc:
            p = min(0.97, p + mej); j = lambda: random.uniform(-0.08, 0.08)
            ej.append((anio, per, nom, gr, 10, round(10*min(1,p+0.1+j())), 10, round(10*min(1,p+0.15+j())), 6, round(6*min(1,p+0.1+j())),
                       round(24*min(1,p+j())), round(3*min(1,p+j()),1), round(1+4*min(1,p+j()),1), round(min(0.95,0.5+0.4*p+j()),2), 2,
                       round(max(0.05,0.1+0.45*p+j()),2), random.randint(20,30), 6 if nom[-1] in 'ADG' else '', round(6*min(1,p+0.05)) if nom[-1] in 'ADG' else '',
                       round(5*min(1,p+0.1)), 1, round(2.6+2.2*p,1) if per == 1 else '', random.randint(18,28) if per == 1 else ''))

NROW = 300
for r in range(2, NROW + 2):
    v = ej[r - 2] if r - 2 < len(ej) else None
    for j, (name, w, kind) in enumerate(cols, 1):
        c = R.cell(r, j); c.border = B; c.font = F()
        if kind == 'in':
            c.fill = INP; c.font = F(color='0000FF')
            if v is not None and v[j - 1] != '': c.value = v[j - 1]
    q = lambda col: f'{col}{r}'
    E,Fc,G,Hh,I,Jj,K,Lc,M,N,O,Pp,Q,Rr,S,T,U,V,Wc = [q(x) for x in 'E F G H I J K L M N O P Q R S T U V W'.split()]
    form = {
     'X': f'=IF(N({E})>0,MIN(100,{Fc}/{E}*100),"")',
     'Y': f'=IF(N({G})>0,MIN(100,{Hh}/{G}*100),"")',
     'Z': f'=IF(N({I})>0,MIN(100,{Jj}/{I}*100),"")',
     'AA': f'=IF({K}="","",{K}/24*100)',
     'AB': f'=IF({Lc}="","",{Lc}/3*100)',
     'AC': f'=IF(OR({M}="",AND({N}<>"",N({N})<{KMIN})),"",MAX(0,MIN(100,({M}-1)/4*100)))',
     'AD': f'=IF(OR({Pp}="",N({Q})=0),"",MAX(0,MIN(100,(({Q}/({Q}+{KC}))*{Pp}+({KC}/({Q}+{KC}))*{GM}-0.1)/0.4*100)))',
     'AE': f'=IF(N({Rr})>0,MIN(100,{S}/{Rr}*100),"")',
     'AF': f'=IF(N({U})>0,MIN(100,{T}/({U}*5)*100),"")',
     'AG': f'=IF(OR({V}="",N({Wc})<10),"",({V}-1)/4*100)',
     'AH': f'=IF($C{r}="","",COUNT(X{r}:AG{r}))',
    }
    comps = ['X','Y','Z','AA','AB','AC','AD','AE','AF','AG']; ws_ = ['C1','C2','C3','C4','C5','C6','C7','C8','C9','C10']
    num = '+'.join(f'IF(ISNUMBER({c}{r}),{c}{r}*{W[w]},0)' for c, w in zip(comps, ws_))
    den = '+'.join(f'IF(ISNUMBER({c}{r}),{W[w]},0)' for c, w in zip(comps, ws_))
    form['AI'] = f'=IF(OR($C{r}="",N(AH{r})<{MINF}),"",({num})/({den}))'
    form['AJ'] = f'=IF(AND(ISNUMBER(Y{r}),N(Y{r})<60),69.9,IF(OR(AND(ISNUMBER(AE{r}),N(AE{r})<80),AND(ISNUMBER(AC{r}),N(AC{r})<50,N({O})>=2)),84.9,100))'
    form['AK'] = f'=IF(AI{r}="","",MIN(AI{r},AJ{r}))'
    form['AL'] = (f'=IF($C{r}="","",IF(AK{r}="","Gris — menos de {3} fuentes",'
                  f'IF(AND(AK{r}<{AMAR},{N}<>"",N({N})<{KMIN},N(Y{r})>=60),"Gris — repetir observación",'
                  f'IF(AK{r}>={VERDE},"Verde",IF(AK{r}>={AMAR},"Amarillo","Rojo")))))')
    form['AM'] = (f'=IF($C{r}="","",IF(AND(ISNUMBER(Y{r}),N(Y{r})<60),"Ejecución < 60 %: rojo automático",'
                  f'IF(AND(ISNUMBER(AE{r}),N(AE{r})<80),"PIAR < 80 %: máximo amarillo",'
                  f'IF(AND(ISNUMBER(AC{r}),N(AC{r})<50,N({O})>=2),"Observación < 50: máximo amarillo",'
                  f'IF(AND({N}<>"",N({N})<{KMIN}),"κ < 0,60: observación descartada, repetir","")))))')
    for col, fx in form.items():
        c = R[f'{col}{r}']; c.value = fx
        c.number_format = '0.0'
R.cell(1, 14).comment = Comment('Kappa ponderado entre dos observadores (lo calcula el aplicativo). Si es < 0,60 la observación no cuenta.', 'CENTEC')
R.cell(1, 16).comment = Comment('Promedio de g = (post − pre)/(100 − pre) por estudiante, excluyendo pre > 90.', 'CENTEC')
for col, color in [('Verde', 'D7F0DF'), ('Amarillo', 'FBEFD2'), ('Rojo', 'F7D9D5'), ('Gris', 'E6E9EC')]:
    R.conditional_formatting.add(f'AL2:AL{NROW+1}', FormulaRule(formula=[f'LEFT($AL2,{len(col)})="{col}"'], fill=PatternFill('solid', fgColor=color)))
dv = DataValidation(type='whole', operator='between', formula1='1', formula2='4', showErrorMessage=True, errorTitle='Período', error='El período va de 1 a 4'); R.add_data_validation(dv); dv.add(f'B2:B{NROW+1}')
dv2 = DataValidation(type='decimal', operator='between', formula1='1', formula2='5'); R.add_data_validation(dv2); dv2.add(f'M2:M{NROW+1}'); dv2.add(f'V2:V{NROW+1}')

# ---------- Tablero ----------
T = wb.create_sheet('Tablero')
T['A1'] = 'TABLERO DEL AÑO'; T['A1'].font = Font(name='Arial', size=14, bold=True, color=NAVY)
T['A3'] = 'Año a analizar'; T['B3'] = 2027; T['B3'].fill = INP; T['B3'].font = F(color='0000FF', bold=True)
T['A4'] = 'Período (0 = año completo)'; T['B4'] = 0; T['B4'].fill = INP; T['B4'].font = F(color='0000FF', bold=True)
for a in ('A3', 'A4'): T[a].font = F(bold=True)
RR = lambda col: f'Registro!${col}$2:${col}${NROW+1}'
cond = f'{RR("A")},$B$3,{RR("B")},IF($B$4=0,">0",$B$4)'
T['A6'] = 'ICD institucional'; T['B6'] = f'=IFERROR(AVERAGEIFS({RR("AK")},{cond}),"—")'
T['A7'] = 'Registros en Verde'; T['B7'] = f'=COUNTIFS({cond},{RR("AL")},"Verde")'
T['A8'] = 'Registros en Amarillo'; T['B8'] = f'=COUNTIFS({cond},{RR("AL")},"Amarillo")'
T['A9'] = 'Registros en Rojo'; T['B9'] = f'=COUNTIFS({cond},{RR("AL")},"Rojo")'
T['A10'] = 'Registros en Gris (faltan datos)'; T['B10'] = f'=COUNTIFS({cond},{RR("AL")},"Gris*")'
for i in range(6, 11): T[f'A{i}'].font = F(bold=True); T[f'B{i}'].font = F(bold=True); T[f'B{i}'].number_format = '0.0'
head(T, 12, ['Componente (promedio institucional)', 'Valor', 'Meta / lectura'], [40, 12, 50])
names = [c['codigo'] + ' ' + c['nombre'] for c in IND['componentes']]
for i, (nm, col) in enumerate(zip(names, ['X','Y','Z','AA','AB','AC','AD','AE','AF','AG']), 13):
    T.cell(i, 1, nm).font = F(); c = T.cell(i, 2, f'=IFERROR(AVERAGEIFS({RR(col)},{cond}),"—")'); c.number_format = '0.0'; c.font = F()
    T.cell(i, 3, '≥ 85 verde · 70–84 amarillo · < 70 rojo').font = F(color='666666')
docs = sorted(set(e[2] for e in ej))
head(T, 25, ['Docente', 'ICD del año/período', 'Estado'])
T['D25'] = 'Escriba o pegue los nombres de los docentes en la columna A'; T['D25'].font = F(italic=True, color='666666')
for i in range(26, 66):
    d = docs[i - 26] if i - 26 < len(docs) else None
    c = T.cell(i, 1, d); c.fill = INP; c.font = F(color='0000FF'); c.border = B
    T.cell(i, 2, f'=IF(A{i}="","",IFERROR(AVERAGEIFS({RR("AK")},{cond},{RR("C")},A{i}),"Sin datos"))').number_format = '0.0'
    T.cell(i, 3, f'=IF(OR(A{i}="",NOT(ISNUMBER(B{i}))),"",IF(B{i}>={VERDE},"Verde",IF(B{i}>={AMAR},"Amarillo","Rojo")))')
    for j in (2, 3): T.cell(i, j).border = B; T.cell(i, j).font = F()
for col, color in [('Verde', 'D7F0DF'), ('Amarillo', 'FBEFD2'), ('Rojo', 'F7D9D5')]:
    T.conditional_formatting.add('C26:C65', FormulaRule(formula=[f'$C26="{col}"'], fill=PatternFill('solid', fgColor=color)))

# ---------- Año vs año ----------
A = wb.create_sheet('Año vs año')
A['A1'] = 'COMPARATIVO INTERANUAL (mismas fórmulas, misma escala)'; A['A1'].font = Font(name='Arial', size=14, bold=True, color=NAVY)
A['A2'] = 'Una variación ≥ 5 puntos se considera cambio real. Agregue años en la fila 4 (texto).'; A['A2'].font = F(italic=True)
years = ['2027', '2028', '2029', '2030', '2031']
head(A, 4, ['Indicador'] + years + ['Δ último año con dato', 'Tendencia'], [40, 10, 10, 10, 10, 10, 16, 14])
rows = [('ICD institucional', 'AK')] + list(zip(names, ['X','Y','Z','AA','AB','AC','AD','AE','AF','AG']))
for i, (nm, col) in enumerate(rows, 5):
    A.cell(i, 1, nm).font = F(bold=(i == 5))
    for j, y in enumerate(years, 2):
        c = A.cell(i, j, f'=IFERROR(AVERAGEIFS({RR(col)},{RR("A")},VALUE({L(j)}$4)),"")'); c.number_format = '0.0'; c.font = F(); c.border = B
    # último par de años con dato
    A.cell(i, 7, f'=IFERROR(LOOKUP(2,1/ISNUMBER(B{i}:F{i}),B{i}:F{i})-INDEX(B{i}:F{i},MATCH(2,1/ISNUMBER(B{i}:F{i}))-1),"")').number_format = '+0.0;-0.0;0.0'
    A.cell(i, 8, f'=IF(G{i}="","—",IF(G{i}>=5,"↑ mejora",IF(G{i}<=-5,"↓ retroceso","= estable")))')
    for j in (7, 8): A.cell(i, j).font = F(); A.cell(i, j).border = B
A.cell(17, 1, 'Docentes en rojo (registros)').font = F()
for j, y in enumerate(years, 2):
    A.cell(17, j, f'=COUNTIFS({RR("A")},VALUE({L(j)}$4),{RR("AL")},"Rojo")').font = F()

# ---------- Lecciones ----------
Le = wb.create_sheet('Lecciones')
head(Le, 1, ['Año', 'Indicador', 'Meta', 'Resultado', 'Brecha', 'Causa raíz (5 porqués)', 'Decisión', 'Acción para el año siguiente', 'Responsable', 'Fecha de verificación'], [7, 30, 9, 10, 9, 45, 12, 45, 18, 14])
dvd = DataValidation(type='list', formula1='"Mantener,Ajustar,Eliminar,Escalar"'); Le.add_data_validation(dvd); dvd.add('G2:G200')
Le.append([2027, 'EJEMPLO: C2 Ejecución', 90, 84, None, 'Semanas de evaluación y eventos no contemplados en la malla', 'Ajustar', 'Reservar 1 semana de margen por período en la malla 2028', 'Coordinación académica', '2028-04-15'])
for r in range(2, 201):
    Le[f'E{r}'] = f'=IF(OR(C{r}="",D{r}=""),"",D{r}-C{r})'
    for j in range(1, 11):
        c = Le.cell(r, j); c.border = B; c.font = F(color='0000FF') if j != 5 else F(); c.alignment = Alignment(wrap_text=True, vertical='top')
        if j != 5: c.fill = INP
# ---------- DASHBOARD (primera hoja, lectura fácil) ----------
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
DB = wb.create_sheet('Dashboard', 0)
wb.active = 0
DB.sheet_view.showGridLines = False
for col, w in zip('ABCDEFGHIJKLMNOPQRSTU', [2] + [11]*12 + [3, 3, 30, 10, 12, 10, 10, 10]):
    DB.column_dimensions[col].width = w
G, Y, RD, GR = '2E9E5B', 'E8A317', 'D2463A', '9AA5AD'
GL, YL, RL, GRL = 'DDF2E5', 'FCEFCF', 'F8DCD8', 'E9ECEE'
big = lambda sz, color='1A1A1A', bold=True: Font(name='Arial', size=sz, bold=bold, color=color)
center = Alignment(horizontal='center', vertical='center', wrap_text=True)
left = Alignment(horizontal='left', vertical='center', wrap_text=True)
def box(rng, fill=None, text=None, font=None, align=center):
    DB.merge_cells(rng); c = DB[rng.split(':')[0]]
    if text is not None: c.value = text
    if font: c.font = font
    c.alignment = align
    if fill:
        for row in DB[rng]:
            for x in row: x.fill = PatternFill('solid', fgColor=fill)
    return c
box('B1:M1', None, '¿CÓMO VAMOS CON LA PLANEACIÓN?', big(22, NAVY))
box('B2:M2', None, 'Tablero de lectura fácil · CENTEC Córdoba · Los números salen solos de la hoja «Registro»', big(11, '555555', False))
DB['B4'] = 'Año que quiere ver:'; DB['B4'].font = big(11); DB.merge_cells('B4:D4')
DB['E4'] = 2027; DB['E4'].fill = INP; DB['E4'].font = big(14, '0000FF'); DB['E4'].alignment = center
DB['G4'] = 'Período (escriba 0 para todo el año):'; DB['G4'].font = big(11); DB.merge_cells('G4:J4')
DB['K4'] = 0; DB['K4'].fill = INP; DB['K4'].font = big(14, '0000FF'); DB['K4'].alignment = center
DB['L4'] = '← cambie solo las casillas amarillas'; DB['L4'].font = Font(name='Arial', size=9, italic=True, color='777777'); DB.merge_cells('L4:M4')
T['B3'] = '=Dashboard!$E$4'; T['B4'] = '=Dashboard!$K$4'; T['B3'].fill = PatternFill(fill_type=None); T['B4'].fill = PatternFill(fill_type=None)
T['C3'] = '← se cambia en la hoja Dashboard'; T['C3'].font = F(italic=True, color='777777')
# Nota del colegio
box('B6:E6', NAVY, 'NOTA DEL COLEGIO (de 0 a 100)', Font(name='Arial', size=11, bold=True, color='FFFFFF'))
c = box('B7:E11', 'F4F6FA', '=IF(ISNUMBER(Tablero!B6),ROUND(Tablero!B6,0),"Sin datos")', big(48, NAVY)); c.number_format = '0'
msg = ('=IF(NOT(ISNUMBER(Tablero!B6)),"Todavía no hay datos suficientes para este año. Registre evidencias en la hoja Registro.",'
       'IF(Tablero!B6>=85,"VERDE. El colegio cumple muy bien lo que planeó. Hay que mantenerlo y compartir lo que funciona.",'
       'IF(Tablero!B6>=70,"AMARILLO. El colegio cumple, pero le falta. Revise abajo qué es lo que más hay que mejorar.",'
       '"ROJO. Lo planeado no se está cumpliendo. Se necesita un plan de mejora inmediato.")))')
box('F6:M11', 'F4F6FA', msg, big(15, '1A1A1A'), left)
for col, color in [('VERDE', GL), ('AMARILLO', YL), ('ROJO', RL)]:
    DB.conditional_formatting.add('F6:M11', FormulaRule(formula=[f'LEFT($F$6,{len(col)})="{col}"'], fill=PatternFill('solid', fgColor=color)))
    DB.conditional_formatting.add('B7:E11', FormulaRule(formula=[f'LEFT($F$6,{len(col)})="{col}"'], fill=PatternFill('solid', fgColor=color)))
# Semáforo en 4 tarjetas
box('B13:M13', None, '¿CUÁNTOS DOCENTES ESTÁN EN CADA COLOR?', big(13, NAVY), left)
cards = [('B', 'D', G, GL, 'Verde', 'cumplen muy bien'), ('E', 'G', Y, YL, 'Amarillo', 'cumplen a medias'), ('H', 'J', RD, RL, 'Rojo', 'no cumplen'), ('K', 'M', GR, GRL, 'Gris', 'faltan datos')]
for c1, c2, col, light, nombre, txt in cards:
    box(f'{c1}14:{c2}14', col, nombre.upper(), Font(name='Arial', size=11, bold=True, color='FFFFFF'))
    cnt = f'=COUNTIF(Tablero!$C$26:$C$65,"{nombre}")' if nombre != 'Gris' else '=COUNTIFS(Tablero!$A$26:$A$65,"<>",Tablero!$B$26:$B$65,"Sin datos")'
    box(f'{c1}15:{c2}17', light, cnt, big(34, col))
    box(f'{c1}18:{c2}18', light, txt, big(10, '333333', False))
box('B19:M19', None, '=IFERROR("De cada 10 docentes, "&ROUND(10*COUNTIF(Tablero!$C$26:$C$65,"Verde")/COUNTA(Tablero!$A$26:$A$65),0)&" están en verde y "&ROUND(10*COUNTIF(Tablero!$C$26:$C$65,"Rojo")/COUNTA(Tablero!$A$26:$A$65),0)&" están en rojo.","")', big(12, '333333', False), left)

# ---- Datos de apoyo para los gráficos (columna P en adelante, visibles) ----
DB['P6'] = 'DATOS DE LOS GRÁFICOS (no editar)'; DB['P6'].font = big(10, '777777')
for j, h in enumerate(['Docente', 'Verde', 'Amarillo', 'Rojo'], 16):
    x = DB.cell(7, j, h); x.font = H; x.fill = HF
ND = 30
for i in range(ND):
    r = 8 + i; tr = 26 + i
    DB.cell(r, 16, f'=IF(Tablero!A{tr}="","",SUBSTITUTE(Tablero!A{tr},"EJEMPLO ",""))')
    for j, nombre in zip((17, 18, 19), ('Verde', 'Amarillo', 'Rojo')):
        x = DB.cell(r, j, f'=IF(Tablero!C{tr}="{nombre}",ROUND(Tablero!B{tr},0),0)'); x.number_format = '0;;'
    for j in range(16, 20): DB.cell(r, j).font = F(color='777777')
# conteo para la dona
for j, h in enumerate(['Estado', 'Docentes'], 21):
    x = DB.cell(7, j, h); x.font = H; x.fill = HF
DB.column_dimensions['U'].width = 12; DB.column_dimensions['V'].width = 10
for i, (nombre, ref) in enumerate([('Verde', 'B15'), ('Amarillo', 'E15'), ('Rojo', 'H15'), ('Gris', 'K15')]):
    DB.cell(8 + i, 21, nombre); DB.cell(8 + i, 22, f'={ref}')
# componentes
short = ['Planeó a tiempo', 'Cumplió las semanas', 'Enseñó los DBA', 'Calidad del planeador', 'Cuadernos', 'Clase observada', 'Lo que aprendieron', 'Inclusión (PIAR)', 'Proyectos transversales', 'Opinión de estudiantes']
DB['U14'] = 'Componente'; DB['V14'] = 'Verde'; DB['W14'] = 'Amarillo'; DB['X14'] = 'Rojo'; DB['Y14'] = 'Valor'
for x in ('U14', 'V14', 'W14', 'X14', 'Y14'): DB[x].font = H; DB[x].fill = HF
DB.column_dimensions['W'].width = 10; DB.column_dimensions['X'].width = 10; DB.column_dimensions['Y'].width = 10
for i, nm in enumerate(short):
    r = 15 + i; tr = 13 + i
    DB.cell(r, 21, nm); DB.cell(r, 25, f'=IF(ISNUMBER(Tablero!B{tr}),ROUND(Tablero!B{tr},0),"")')
    DB.cell(r, 22, f'=IF(AND(ISNUMBER(Y{r}),N(Y{r})>=85),Y{r},0)').number_format = '0;;'
    DB.cell(r, 23, f'=IF(AND(ISNUMBER(Y{r}),N(Y{r})>=70,N(Y{r})<85),Y{r},0)').number_format = '0;;'
    DB.cell(r, 24, f'=IF(AND(ISNUMBER(Y{r}),N(Y{r})<70),Y{r},0)').number_format = '0;;'
    for j in range(21, 26): DB.cell(r, j).font = F(color='777777')
# años
DB['U27'] = 'Año'; DB['V27'] = 'Nota del colegio'
for x in ('U27', 'V27'): DB[x].font = H; DB[x].fill = HF
for i in range(5):
    col = 'BCDEF'[i]
    DB.cell(28 + i, 21, f"='Año vs año'!{col}4"); x = DB.cell(28 + i, 22, f"=IF(ISNUMBER('Año vs año'!{col}5),ROUND('Año vs año'!{col}5,0),0)"); x.number_format = '0;;'

def estilo_barras(ch, cats, data_ref, colores, titulo, alto, ancho, horizontal=True, stacked=True):
    ch.type = 'bar' if horizontal else 'col'
    if stacked: ch.grouping = 'stacked'; ch.overlap = 100
    ch.title = titulo; ch.style = 10; ch.height = alto; ch.width = ancho
    ch.add_data(data_ref, titles_from_data=True); ch.set_categories(cats)
    for sr, col in zip(ch.series, colores):
        sr.graphicalProperties.solidFill = col; sr.graphicalProperties.line.solidFill = col
        sr.dLbls = DataLabelList(); sr.dLbls.showVal = True; sr.dLbls.numFmt = '0;;'
        sr.dLbls.showSerName = False; sr.dLbls.showCatName = False; sr.dLbls.showLegendKey = False
    ch.y_axis.scaling.min = 0; ch.y_axis.scaling.max = 100; ch.y_axis.majorGridlines = None
    ch.y_axis.delete = True; ch.x_axis.delete = False; ch.legend.position = 'b'
    if horizontal: ch.x_axis.scaling.orientation = 'maxMin'
    ch.gapWidth = 40
    return ch

box('B21:M21', None, 'NOTA DE CADA DOCENTE (de 0 a 100) · el color dice cómo va', big(13, NAVY), left)
ch1 = estilo_barras(BarChart(), Reference(DB, min_col=16, min_row=8, max_row=8 + 11), Reference(DB, min_col=17, max_col=19, min_row=7, max_row=8 + 11), [G, Y, RD], None, 10, 17)
DB.add_chart(ch1, 'B23')
box('B43:M43', None, '=IF(COUNTA(Tablero!A26:A65)>12,"Hay más de 12 docentes: el gráfico muestra los 12 primeros. La lista completa está en la hoja Tablero.","El gráfico muestra hasta 12 docentes; la lista completa está en la hoja Tablero.")', Font(name='Arial', size=9, italic=True, color='777777'), left)

box('B45:F45', None, 'DOCENTES POR COLOR', big(13, NAVY), left)
dn = DoughnutChart(); dn.add_data(Reference(DB, min_col=22, min_row=7, max_row=11), titles_from_data=True); dn.set_categories(Reference(DB, min_col=21, min_row=8, max_row=11))
dn.holeSize = 55; dn.height = 8; dn.width = 9; dn.title = None
for k, col in enumerate([G, Y, RD, GR]):
    pt = DataPoint(idx=k); pt.graphicalProperties.solidFill = col; dn.series[0].dPt.append(pt)
dn.series[0].dLbls = DataLabelList(); dn.series[0].dLbls.showVal = True; dn.series[0].dLbls.showCatName = False; dn.series[0].dLbls.showSerName = False; dn.series[0].dLbls.showLegendKey = False; dn.series[0].dLbls.showPercent = False; dn.series[0].dLbls.numFmt = '0;;'
dn.legend.position = 'r'
DB.add_chart(dn, 'B47')

box('H44:M44', None, '¿EN QUÉ SOMOS BUENOS Y QUÉ FALTA?', big(13, NAVY), left)
vals = 'Dashboard!$Y$15:$Y$24'; nms = 'Dashboard!$U$15:$U$24'
box('H45:M46', GL, f'=IFERROR("Lo que mejor hacemos: "&INDEX({nms},MATCH(MAX({vals}),{vals},0))&" ("&MAX({vals})&" de 100)","Sin datos")', big(12, '1B5E35'), left)
box('H47:M48', RL, f'=IFERROR("Lo que más debemos mejorar: "&INDEX({nms},MATCH(MIN({vals}),{vals},0))&" ("&MIN({vals})&" de 100)","Sin datos")', big(12, '8A2A20'), left)
box('H49:M50', YL, '=IFERROR("Componentes por debajo de 70: "&COUNTIF(Dashboard!$Y$15:$Y$24,"<70")&" de "&COUNT(Dashboard!$Y$15:$Y$24),"")', big(12, '6B4A00'), left)
box('H51:M53', 'F4F6FA', 'Meta de cada componente: 85 o más (verde). Entre 70 y 84 es amarillo. Menos de 70 es rojo.', big(10, '333333', False), left)

box('B64:M64', None, '¿QUÉ PARTE DEL TRABAJO VA MEJOR Y CUÁL PEOR? (promedio del colegio, de 0 a 100)', big(13, NAVY), left)
ch2 = estilo_barras(BarChart(), Reference(DB, min_col=21, min_row=15, max_row=24), Reference(DB, min_col=22, max_col=24, min_row=14, max_row=24), [G, Y, RD], None, 9, 17)
DB.add_chart(ch2, 'B66')

box('B86:M86', None, '¿VAMOS MEJORANDO AÑO TRAS AÑO? (nota del colegio)', big(13, NAVY), left)
ch3 = BarChart(); ch3.type = 'col'; ch3.style = 10; ch3.height = 8; ch3.width = 17; ch3.title = None
ch3.add_data(Reference(DB, min_col=22, min_row=27, max_row=32), titles_from_data=True); ch3.set_categories(Reference(DB, min_col=21, min_row=28, max_row=32))
ch3.series[0].graphicalProperties.solidFill = NAVY
ch3.series[0].dLbls = DataLabelList(); ch3.series[0].dLbls.showVal = True; ch3.series[0].dLbls.numFmt = '0;;'; ch3.series[0].dLbls.showSerName = False; ch3.series[0].dLbls.showCatName = False; ch3.series[0].dLbls.showLegendKey = False
ch3.y_axis.scaling.min = 0; ch3.y_axis.scaling.max = 100; ch3.y_axis.delete = True; ch3.y_axis.majorGridlines = None; ch3.legend = None; ch3.gapWidth = 60
DB.add_chart(ch3, 'B88')
box('B104:M105', 'F4F6FA', "=IF('Año vs año'!H5=\"—\",\"Cuando haya dos años con datos, aquí verá si el colegio mejoró, se mantuvo o retrocedió.\",\"Comparado con el año anterior, el colegio \"&IF('Año vs año'!G5>=5,\"MEJORÓ\",IF('Año vs año'!G5<=-5,\"RETROCEDIÓ\",\"SE MANTUVO IGUAL\"))&\" (\"&TEXT('Año vs año'!G5,\"+0;-0;0\")&\" puntos).\")", big(13, NAVY), left)

box('B107:M107', None, 'CÓMO LEER ESTE TABLERO', big(13, NAVY), left)
guia = ['La nota va de 0 a 100. Más alta es mejor.',
        'Verde (85 o más): cumple muy bien. Amarillo (70 a 84): cumple a medias. Rojo (menos de 70): no cumple. Gris: faltan datos para calificar.',
        'La nota no la pone una sola persona: combina planeadores, semanas cumplidas, cuadernos sorteados, clases observadas por dos personas, lo que aprendieron los estudiantes y su opinión.',
        'Una diferencia de 5 puntos o más entre años es un cambio real; menos de 5 se considera igual.',
        'Para ver otro año o período, cambie las casillas amarillas de arriba.']
for i, t in enumerate(guia):
    box(f'B{108+i}:M{108+i}', None, '•  ' + t, big(11, '333333', False), left)
    DB.row_dimensions[108 + i].height = 30
DB.row_dimensions[1].height = 34; DB.row_dimensions[19].height = 22
for r in range(7, 12): DB.row_dimensions[r].height = 22
DB.freeze_panes = 'A5'

wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT); print('ok', OUT)
